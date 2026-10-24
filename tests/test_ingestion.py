"""Tests for loaders (text normalization, DOCX table order) and chunking (offsets, determinism)."""

from pathlib import Path

import docx

from rag_platform.config import RagConfig
from rag_platform.ingestion.chunking import chunk_text
from rag_platform.ingestion.loaders import load_corpus, load_document

CFG = RagConfig(chunk_size=120, chunk_overlap=20)


def test_load_text_and_markdown(tmp_path: Path):
    (tmp_path / "a.md").write_bytes(b"# Title\r\n\r\nHello world")
    (tmp_path / "b.txt").write_text("plain", encoding="utf-8")
    (tmp_path / "skip.bin").write_bytes(b"x")
    docs = load_corpus(tmp_path)
    assert [d.path.name for d in docs] == ["a.md", "b.txt"]
    assert docs[0].text == "# Title\n\nHello world\n"


def test_load_docx_keeps_table_in_order(tmp_path: Path):
    d = docx.Document()
    d.add_paragraph("Before")
    t = d.add_table(rows=2, cols=2)
    for r, row in enumerate([["Plan", "Limit"], ["Basic", "10"]]):
        for c, v in enumerate(row):
            t.cell(r, c).text = v
    d.add_paragraph("After")
    d.save(tmp_path / "x.docx")
    text = load_document(tmp_path / "x.docx").text
    assert text.index("Before") < text.index("Basic | 10") < text.index("After")


def test_chunk_offsets_point_into_source_text():
    text = ("Sentence number one is here. " * 30 + "\n\n") * 3
    chunks = chunk_text(text, CFG)
    assert len(chunks) > 3
    for c in chunks:
        assert text[c.start : c.end] == c.text
        assert len(c.text) <= CFG.chunk_size
    assert len({c.text_hash for c in chunks}) == len({c.text for c in chunks})


def test_chunking_is_deterministic():
    text = "alpha beta gamma. " * 100
    assert chunk_text(text, CFG) == chunk_text(text, CFG)
