"""Generate the DOCX and PDF files of the Lumora Labs AI Engineering Handbook sample corpus.

MD files in data/sample_corpus/ are static. Run: uv run python scripts/make_corpus.py

Also writes data/corpus_updates/model_selection_policy_v1.docx (policy version 2.0, same file
name so copying it over the original tests incremental re-indexing of a changed document).
"""

from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import (
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

ROOT = Path(__file__).resolve().parent.parent
CORPUS = ROOT / "data" / "sample_corpus"
UPDATES = ROOT / "data" / "corpus_updates"

# Blocks: ("h", text) | ("p", text) | ("l", [items]) | ("t", rows with rows[0] = header)
#         | ("pb", None) page break

# --------------------------------------------------------------------------------------------
# ai_engineering_handbook.docx
# --------------------------------------------------------------------------------------------

HANDBOOK = [
    (
        "p",
        "AI Engineering Handbook - version 3.2 - effective 2026-03-15. Owner: Nadia Ferreira, "
        "Head of AI Engineering.",
    ),
    ("h", "1. About the team"),
    (
        "p",
        "The AI Engineering team at Lumora Labs builds and operates the company's LLM-based "
        "features: retrieval-augmented assistants, document processing pipelines and agentic "
        "workflows. The team has 18 engineers organised in two squads: the Retrieval squad and "
        "the Agents squad. A shared ML Platform group runs the model gateway, the vector "
        "infrastructure and the evaluation pipeline.",
    ),
    ("h", "2. Roles and responsibilities"),
    (
        "t",
        [
            ["Role", "Person", "Responsibilities"],
            [
                "Head of AI Engineering",
                "Nadia Ferreira",
                "Team direction, chairs the Model Review Board, incident escalation for AI-SEV1",
            ],
            [
                "Retrieval squad tech lead",
                "Mateo Lindqvist",
                "RAG architecture, owner of the RAG Design Guide",
            ],
            [
                "Agents squad tech lead",
                "Ravi Menon",
                "Agent architecture, owner of the Agent Design Patterns guide",
            ],
            [
                "Evaluation Lead",
                "Tomasz Wierzbicki",
                "Golden sets, LLM-as-judge, release gates, owner of the Evaluation Playbook",
            ],
            [
                "AI Safety & Governance Lead",
                "Keiko Arai",
                "Guardrails, red-teaming, owner of the LLMOps Incident Runbook",
            ],
            [
                "ML Platform lead",
                "Lena Okafor",
                "Lumora Gateway, vector infrastructure, embedding service, cost dashboards",
            ],
        ],
    ),
    (
        "p",
        "The Model Review Board (MRB) approves the use of models for new use cases and any "
        "spend above the caps in the Model Selection Policy. It is chaired by the Head of AI "
        "Engineering and meets every second Thursday.",
    ),
    ("h", "3. LLM fundamentals"),
    (
        "p",
        "A large language model (LLM) is a neural network, almost always a transformer, trained "
        "to predict the next token in a sequence. Tokens are sub-word units produced by a "
        "tokenizer; in English text one token is on average about three quarters of a word, so "
        "token counts, not word counts, determine cost and context usage.",
    ),
    (
        "p",
        "The context window is the maximum number of tokens the model can attend to in one "
        "request, counting both the input (system prompt, instructions, retrieved documents, "
        "conversation history) and the generated output. Content beyond the window is not seen "
        "at all, and very long contexts can reduce how reliably the model uses information "
        "placed in the middle.",
    ),
    (
        "p",
        "Generation is a sampling process. Temperature rescales the probability distribution "
        "over the next token: low values make output more deterministic, high values make it "
        "more varied. Top-p (nucleus) sampling restricts sampling to the smallest set of tokens "
        "whose cumulative probability exceeds p. Even at temperature 0, outputs from hosted "
        "models are not guaranteed to be bit-for-bit reproducible.",
    ),
    (
        "p",
        "A hallucination is output that is fluent but not supported by the input or by fact. "
        "LLMs have no built-in notion of truth; they produce plausible continuations. Grounding "
        "answers in retrieved documents, requiring citations and allowing the model to abstain "
        "are the main mitigations used by the team.",
    ),
    (
        "p",
        "An embedding is a fixed-length vector representation of a text, produced by an "
        "embedding model, such that semantically similar texts have nearby vectors. Embeddings "
        "power semantic search and are described in detail in the RAG Design Guide.",
    ),
    ("h", "4. Prompt engineering standards"),
    (
        "l",
        [
            "Prompts are code. Every production prompt lives in the repository under the "
            "prompts/ directory and is reviewed like any other change.",
            "Every prompt has an identifier of the form PRM-<squad>-<number>, for example "
            "PRM-RET-014, and a version number that is logged with every model call.",
            "The system prompt must state the role, the task, the allowed sources of "
            "information and the exact output format.",
            "Use structured output (a JSON schema) whenever the result is consumed by code.",
            "Default temperature is 0.2 for extraction, classification and question answering. "
            "Creative tasks may use up to 0.7. Temperatures above 1.0 are not allowed in "
            "production.",
            "Few-shot prompts may contain at most 5 examples; examples must not contain real "
            "customer data.",
            "Delimit untrusted content (user input, retrieved documents, tool output) clearly "
            "and instruct the model to treat it as data, not instructions.",
            "A prompt change must be evaluated on the golden set before merge. No release-gate "
            "metric may drop by more than 2 points compared with the current production version.",
        ],
    ),
    ("h", "5. Calling models"),
    (
        "p",
        "All model calls go through the Lumora Gateway, run by the ML Platform group. The "
        "gateway enforces authentication, per-feature budgets and logging. Client code must set "
        "a timeout of 30 seconds and may retry at most 2 times with exponential backoff. Direct "
        "calls to model providers from application code are not allowed.",
    ),
    ("h", "6. Code review rules"),
    (
        "l",
        [
            "Every pull request needs at least one approval. Changes to prompts, retrieval "
            "configuration or agent tools need two approvals, one of them from a squad tech lead.",
            "Pull requests that change prompts or retrieval configuration must link the "
            "evaluation report produced by the CI evaluation job.",
            "Pull requests should stay under 400 changed lines, excluding generated files and "
            "test fixtures.",
            "Secrets such as API keys must never be committed; the secret scanner blocks the "
            "merge if one is detected.",
            "Unit tests must not call hosted models. Use fake models and fake embedders so tests "
            "are deterministic and free.",
            "Deterministic logic (parsing, scoring, formatting) is written in plain code and "
            "unit-tested; it is never delegated to a model.",
        ],
    ),
]

# --------------------------------------------------------------------------------------------
# model_selection_policy_v1.docx (v1 in the corpus, v2 in corpus_updates)
# --------------------------------------------------------------------------------------------


def model_selection_policy(v: dict) -> list:
    rows = [["Use case", "Approved model tier", "Max cost per 1K requests", "Approval owner"]]
    rows += [list(r) for r in v["rows"]]
    return [
        (
            "p",
            f"Model Selection Policy - version {v['ver']} - effective {v['date']}. "
            "Owner: Model Review Board (chair: Nadia Ferreira).",
        ),
        ("h", "1. Purpose"),
        (
            "p",
            "This policy defines which class of model may be used for each type of use case, "
            "the maximum spend per 1,000 requests, and who approves exceptions. It refers to "
            "models by tier, not by vendor or product name, so that the policy stays valid when "
            "individual models are replaced.",
        ),
        ("h", "2. Model tiers"),
        (
            "t",
            [
                ["Tier", "Description", "Hosting"],
                ["Tier-1", "Largest, most capable hosted LLM", "Hosted by an external provider"],
                ["Tier-2", "Mid-size hosted LLM, lower cost and latency", "Hosted externally"],
                ["Tier-3", "Small open-weight model", "Self-hosted on Lumora infrastructure"],
            ],
        ),
        ("h", "3. Approved tier per use case"),
        ("t", rows),
        (
            "p",
            "Costs are measured in US dollars at the Lumora Gateway and include input and output "
            "tokens. A use case not listed in the table requires approval by the Model Review "
            "Board before any production traffic.",
        ),
        ("h", "4. Rules"),
        (
            "l",
            [
                "A lower tier than the approved tier may always be used if the evaluation release "
                "gates are still met.",
                "Using a higher tier, or exceeding the cost cap, requires approval by the Model "
                "Review Board.",
                "Data classified as Confidential-Restricted may only be processed by Tier-3 "
                "self-hosted models.",
                f"This policy is reviewed {v['review']}.",
            ]
            + v["extra_rules"],
        ),
    ]


MODEL_POLICY_V1 = model_selection_policy(
    {
        "ver": "1.0",
        "date": "2026-02-01",
        "review": "every 6 months",
        "rows": [
            ("Customer-facing chat assistant", "Tier-1", "USD 40.00", "Model Review Board"),
            ("Internal knowledge Q&A (RAG)", "Tier-2", "USD 12.00", "Head of AI Engineering"),
            ("Document summarisation", "Tier-2", "USD 8.00", "Squad tech lead"),
            ("Document classification", "Tier-3", "USD 1.50", "Squad tech lead"),
            ("Code assistance", "Tier-1", "USD 25.00", "Head of AI Engineering"),
            ("LLM-as-judge evaluation", "Tier-1", "USD 30.00", "Evaluation Lead"),
            ("Embeddings", "Tier-3", "USD 0.20", "ML Platform lead"),
        ],
        "extra_rules": [],
    }
)

MODEL_POLICY_V2 = model_selection_policy(
    {
        "ver": "2.0",
        "date": "2026-09-01",
        "review": "every quarter",
        "rows": [
            ("Customer-facing chat assistant", "Tier-1", "USD 35.00", "Model Review Board"),
            ("Internal knowledge Q&A (RAG)", "Tier-2", "USD 10.00", "Head of AI Engineering"),
            ("Document summarisation", "Tier-3", "USD 3.00", "Squad tech lead"),
            ("Document classification", "Tier-3", "USD 1.00", "Squad tech lead"),
            ("Code assistance", "Tier-1", "USD 25.00", "Head of AI Engineering"),
            ("LLM-as-judge evaluation", "Tier-1", "USD 30.00", "Evaluation Lead"),
            ("Embeddings", "Tier-3", "USD 0.20", "ML Platform lead"),
            ("Agentic workflows", "Tier-1", "USD 60.00", "Model Review Board"),
        ],
        "extra_rules": [
            "Agentic workflows must also follow the limits in the Agent Design Patterns guide.",
        ],
    }
)

# --------------------------------------------------------------------------------------------
# rag_design_guide.pdf (~12 pages)
# --------------------------------------------------------------------------------------------

RAG_GUIDE = [
    (
        "p",
        "RAG Design Guide - version 2.1 - effective 2026-06-10. Owner: Retrieval squad "
        "(tech lead: Mateo Lindqvist). Applies to every retrieval-augmented feature built by "
        "the Lumora Labs AI Engineering team.",
    ),
    ("h", "1. What retrieval-augmented generation is"),
    (
        "p",
        "Retrieval-augmented generation (RAG) answers a question in two stages. First, a "
        "retriever searches a document collection for the passages most relevant to the "
        "question. Second, a language model generates an answer using those passages as "
        "context. The model's parametric knowledge (what it learned during training) is "
        "complemented by non-parametric knowledge (the documents), which can be updated at any "
        "time without retraining.",
    ),
    (
        "p",
        "RAG addresses three weaknesses of a plain LLM: it does not know private or recent "
        "information, it cannot show where an answer came from, and it tends to produce "
        "plausible but unsupported statements. A well-built RAG system answers from the "
        "documents, cites the passages it used, and says that it does not know when the "
        "documents do not contain the answer.",
    ),
    (
        "p",
        "A RAG system has an offline indexing pipeline and an online query pipeline. The "
        "indexing pipeline loads documents, splits them into chunks, computes an embedding for "
        "each chunk, and stores chunks and vectors. The query pipeline embeds the question, "
        "retrieves candidate chunks, optionally reranks them, assembles a prompt, generates an "
        "answer, and validates citations. Most quality problems in production RAG systems are "
        "retrieval problems, not generation problems: if the right passage is not retrieved, "
        "no prompt can fix the answer.",
    ),
    (
        "p",
        "This guide describes each stage, the trade-offs involved, and the defaults the team "
        "has agreed on. Defaults are collected in the table in section 11. Deviating from a "
        "default is allowed when an evaluation report shows the change improves the release "
        "gate metrics defined in the LLM Evaluation Playbook.",
    ),
    ("pb", None),
    ("h", "2. Document ingestion and parsing"),
    (
        "p",
        "Retrieval quality is bounded by parsing quality. Each file format needs its own "
        "loader. Plain text and Markdown can be read directly. Word documents should be read "
        "in document order, including tables, so that a table stays next to the paragraph that "
        "introduces it. PDFs are the hardest format: a PDF stores positioned glyphs, not "
        "paragraphs, and text extraction may merge columns, drop table structure or insert "
        "hyphenation and headers on every page. Scanned PDFs require OCR before any of this "
        "applies.",
    ),
    (
        "p",
        "After extraction, text is normalised: line endings are unified, trailing whitespace "
        "is removed and runs of blank lines are collapsed. Normalisation must happen before "
        "chunking, and the normalised text must be stored, because chunk character offsets and "
        "evaluation evidence spans refer to the normalised text and not to the original file.",
    ),
    (
        "p",
        "Every document is stored with a content hash, computed on the normalised text. When "
        "the indexing pipeline runs again, unchanged documents (same hash) are skipped. A "
        "changed document has all of its old chunks removed from the index before the new "
        "chunks are added; updating only the chunks that differ is not worth the complexity. A "
        "document whose source file has been deleted must be removed from the index, otherwise "
        "the assistant keeps citing a policy that no longer exists.",
    ),
    (
        "p",
        "Metadata captured at ingestion time includes the source path, document type, title, "
        "version and effective date where available, and the ingestion timestamp. Metadata is "
        "used for filtering (for example, only current policies) and for showing citations to "
        "users.",
    ),
    ("pb", None),
    ("h", "3. Chunking"),
    (
        "p",
        "Documents are split into chunks because embedding models have input limits, because "
        "a single vector cannot represent a long document well, and because the generator "
        "should receive only the relevant parts of a document. Chunk size is a trade-off. Small "
        "chunks give precise matches but lose surrounding context; large chunks keep context "
        "but dilute the embedding with unrelated content and consume more of the context "
        "window.",
    ),
    (
        "p",
        "Fixed-size chunking cuts text every N characters or tokens. It is simple but cuts "
        "through sentences and tables. Recursive chunking tries a list of separators in order, "
        "for example paragraph breaks, then line breaks, then sentence ends, then spaces, and "
        "only falls back to the next separator when a piece is still too large. It keeps "
        "paragraphs intact whenever they fit. Structure-aware chunking uses the document's own "
        "structure, such as Markdown headings or DOCX heading styles, to split at section "
        "boundaries and attaches the section heading to each chunk. Semantic chunking places "
        "boundaries where the embedding similarity between consecutive sentences drops; it can "
        "produce better boundaries but costs embedding calls at indexing time and is harder to "
        "reproduce.",
    ),
    (
        "p",
        "Overlap repeats the end of one chunk at the start of the next, so that a fact that "
        "falls on a boundary is fully contained in at least one chunk. Overlap increases index "
        "size and can return near-duplicate chunks; the team default is 100 characters, which "
        "is 12.5% of the default chunk size of 800 characters.",
    ),
    (
        "p",
        "Tables need special care. A table row that is separated from its header row loses its "
        "meaning: the row 'G4 | 35' means nothing without the column names. Team rule: tables "
        "must not be split in the middle of a row, and when a table is longer than one chunk, "
        "every chunk that contains table rows repeats the header row.",
    ),
    (
        "p",
        "Every chunk stores the character offsets of its start and end in the normalised "
        "document text and a hash of its text. Offsets allow citations to point to an exact "
        "span and allow evaluation to decide whether a retrieved chunk overlaps the evidence "
        "for a question, independently of how the document was chunked.",
    ),
    ("pb", None),
    ("h", "4. Embeddings"),
    (
        "p",
        "An embedding model maps a text to a dense vector of fixed dimension, typically a few "
        "hundred to a few thousand floating-point numbers. The model is trained so that texts "
        "with similar meaning have vectors that point in similar directions. Similarity is "
        "usually measured with cosine similarity, the cosine of the angle between two vectors. "
        "If vectors are normalised to unit length, cosine similarity equals the inner (dot) "
        "product, which vector indexes compute efficiently. Team rule: all vectors are L2-"
        "normalised before they are stored.",
    ),
    (
        "p",
        "Many embedding models are asymmetric: they expect a different task setting, or a "
        "different prefix, for queries and for documents, because a short question and a long "
        "passage that answers it are different kinds of text. Using the document setting for "
        "queries usually lowers retrieval quality. Check the model card and use the query "
        "setting at query time.",
    ),
    (
        "p",
        "Absolute similarity scores are not comparable across embedding models. Some models "
        "produce scores spread between 0 and 1, others compress almost all scores into a narrow "
        "band, so that unrelated passages score 0.6 and relevant ones 0.7. A similarity "
        "threshold for abstention must therefore be calibrated on the golden set for the "
        "specific model, never copied from another system.",
    ),
    (
        "p",
        "Some models support reduced output dimensions. Smaller vectors reduce memory and "
        "search time at a small cost in quality. The team default is 768 dimensions for "
        "production indexes; experiments may use the model's full dimension.",
    ),
    (
        "p",
        "Embeddings are cached in the source database, keyed by the hash of the chunk text and "
        "the embedding model name. Re-indexing a document whose chunks did not change therefore "
        "costs no embedding calls. When the embedding model changes, every vector must be "
        "recomputed: vectors from two different models are not comparable and must never be "
        "mixed in one index.",
    ),
    ("pb", None),
    ("h", "5. Vector stores and indexes"),
    (
        "p",
        "A vector index answers the question: which stored vectors are closest to this query "
        "vector? An exact (flat) index compares the query with every vector. It is simple, has "
        "perfect recall and is fast enough for up to roughly a million vectors on a single "
        "machine. Approximate nearest neighbour (ANN) indexes trade a little recall for much "
        "faster search on large collections.",
    ),
    (
        "p",
        "HNSW (hierarchical navigable small world) builds a layered proximity graph and "
        "searches it greedily from the top layer down. Its main parameters are M, the number "
        "of links per node, which controls memory and graph quality; efConstruction, the "
        "candidate list size during building; and efSearch, the candidate list size at query "
        "time, which trades latency for recall. IVF (inverted file) indexes cluster the "
        "vectors and search only the clusters nearest to the query; the nprobe parameter sets "
        "how many clusters are searched. Product quantisation (PQ) compresses vectors into "
        "short codes to reduce memory, at a further cost in accuracy.",
    ),
    (
        "p",
        "Team rule: the relational database is the source of truth for documents, chunks, "
        "offsets and cached embeddings. The vector index is a derived artefact keyed by chunk "
        "id and can be rebuilt from the database at any time. When the index and the database "
        "disagree, the index is rebuilt (error code LUM-1002 in the LLMOps Incident Runbook). "
        "Index ids must never be reused for different chunks.",
    ),
    (
        "p",
        "Metadata filtering restricts search to chunks matching conditions such as document "
        "type or effective date. Filtering after search can return fewer than K results; "
        "filtering before or during search is preferred when the vector store supports it.",
    ),
    ("pb", None),
    ("h", "6. Hybrid search"),
    (
        "p",
        "Dense retrieval matches meaning but can miss exact terms: product codes, error codes "
        "such as LUM-2001, names, and rare acronyms. Lexical retrieval with BM25 matches exact "
        "terms well but misses paraphrases. BM25 scores a document by the frequency of the "
        "query terms in it, weighted by how rare each term is across the collection (inverse "
        "document frequency) and normalised by document length. Its parameter k1, usually "
        "between 1.2 and 2.0, controls term-frequency saturation, and b, usually 0.75, controls "
        "length normalisation.",
    ),
    (
        "p",
        "Hybrid search runs both retrievers and merges their result lists. Because BM25 scores "
        "and cosine similarities are on different scales, the team merges by rank, not by "
        "score, using reciprocal rank fusion (RRF). Each document receives the sum over "
        "retrievers of 1 / (k + rank), where rank is its position in that retriever's list and "
        "k is a constant, conventionally 60, that reduces the influence of the very top ranks. "
        "RRF needs no score normalisation and no tuning of weights, which makes it robust.",
    ),
    (
        "p",
        "Team default: hybrid search is enabled for every production index. Each retriever "
        "returns 50 candidates, the lists are fused with RRF using k = 60, and the fused list "
        "is passed to the reranker. Hybrid search is particularly important for corpora that "
        "contain codes, identifiers and tables.",
    ),
    ("pb", None),
    ("h", "7. Reranking"),
    (
        "p",
        "Embedding models are bi-encoders: the query and each document are encoded separately, "
        "which makes search fast because document vectors are computed once, but limits "
        "accuracy because the model never sees query and document together. A cross-encoder "
        "reranker takes the query and one candidate passage as a single input and outputs a "
        "relevance score. It is far more accurate but far too slow to apply to a whole "
        "collection, so it is applied only to the candidates returned by the first stage.",
    ),
    (
        "p",
        "Team default: rerank the top 50 fused candidates with a cross-encoder and keep the "
        "top 5 for the prompt. An LLM can also be used as a reranker or as a relevance grader "
        "that labels each candidate as relevant or not; this is more expensive and is used in "
        "the agentic query pipeline, not in the fast pipeline.",
    ),
    (
        "p",
        "Reranking changes the score distribution. When a reranker is used, the abstention "
        "threshold must be calibrated on reranker scores, not on embedding similarities.",
    ),
    ("pb", None),
    ("h", "8. Context assembly, citations and abstention"),
    (
        "p",
        "The selected chunks are numbered and placed in the prompt with their source, inside "
        "clear delimiters. The instructions tell the model to answer only from the passages, "
        "to cite the passage number for every claim, and to state that the answer is not "
        "available when the passages do not contain it. Retrieved passages are untrusted "
        "content: instructions that appear inside a document must be ignored.",
    ),
    (
        "p",
        "Citations are validated in code after generation. A citation that refers to a passage "
        "number that was not supplied is removed. An answer that claims to be answerable but "
        "contains no valid citation is replaced by an abstention and logged with error code "
        "LUM-3001.",
    ),
    (
        "p",
        "Abstention has two layers. Before generation, if no retrieved chunk reaches the "
        "calibrated similarity threshold, the system abstains without calling the model "
        "(logged as LUM-1001). After generation, the model's structured output includes an "
        "answerable field, and the system abstains when the model reports that the passages "
        "are insufficient. An honest 'I don't know' is always better than a confident wrong "
        "answer; abstention accuracy is a release gate metric.",
    ),
    (
        "p",
        "The total context passed to the generator is capped at 6,000 tokens. Ordering matters "
        "less with five short passages than with long contexts, but the most relevant passage "
        "is placed first.",
    ),
    ("pb", None),
    ("h", "9. Query transformation"),
    (
        "p",
        "Users write short, ambiguous or conversational questions. Query transformation "
        "rewrites the question before retrieval to improve recall. Query rewriting asks a model "
        "to produce a clearer, self-contained search query, for example by resolving pronouns "
        "from the conversation history ('what about its overlap?' becomes 'what is the default "
        "chunk overlap?'). Multi-query retrieval generates several paraphrases, retrieves for "
        "each, and fuses the results, which helps when the user's wording differs from the "
        "documents' wording.",
    ),
    (
        "p",
        "Query decomposition splits a compound question into sub-questions that are retrieved "
        "separately, for example a question that compares two policies. HyDE (hypothetical "
        "document embeddings) asks the model to write a hypothetical answer and embeds that "
        "answer instead of the question, on the idea that an answer-shaped text is closer in "
        "embedding space to the real answer passage than the question is. HyDE can help on "
        "short queries but can also pull retrieval towards the model's own hallucinated "
        "content.",
    ),
    (
        "p",
        "Every transformation adds at least one model call of latency and cost, and a bad "
        "rewrite can make retrieval worse. Team rule: query transformation is not used in the "
        "fast pipeline. In the agentic pipeline, the query is rewritten only when the relevance "
        "grader finds no relevant passage, and at most 2 rewrites are attempted before the "
        "system abstains.",
    ),
    (
        "p",
        "The original user question, not the rewritten query, is always passed to the "
        "generator and recorded in the evaluation log, so that answers are judged against what "
        "the user actually asked.",
    ),
    ("pb", None),
    ("h", "10. Operating and maintaining indexes"),
    (
        "p",
        "An index is not built once. Documents are added, changed and retired, and the "
        "indexing pipeline must keep the index consistent with the corpus. Incremental "
        "indexing relies on the content hash of each document and the embedding cache keyed "
        "by chunk-text hash, so that a typical nightly run re-embeds only the chunks that "
        "actually changed. A full rebuild is needed only when the embedding model or the "
        "chunking configuration changes; in that case the new index is built side by side and "
        "switched over after its evaluation run passes the release gates.",
    ),
    (
        "p",
        "Each document stores the chunking configuration it was split with. If the configured "
        "chunk size or overlap differs from the stored one, the document is treated as changed "
        "and is re-chunked, even when its content is identical. This prevents an index that "
        "silently mixes chunks produced by different configurations.",
    ),
    (
        "p",
        "The ML Platform group monitors the following signals for every production index: the "
        "number of documents and chunks, the age of the oldest unprocessed document change, the "
        "embedding cache hit rate during indexing, the rate of LUM-1001 abstentions at query "
        "time, and the rate of LUM-1002 index consistency errors. A sudden rise in LUM-1001 "
        "abstentions after an indexing run usually means documents failed to parse, for "
        "example a PDF that lost its text layer.",
    ),
    (
        "p",
        "Retired documents must be removed from the index on the same day they are withdrawn. "
        "Superseded policy versions are either deleted or tagged with an end date and excluded "
        "by a metadata filter, so that the assistant never answers from an outdated version "
        "when a newer one exists.",
    ),
    (
        "p",
        "Backups cover the source database only. Because the vector index is a derived "
        "artefact, it is rebuilt from the restored database rather than backed up separately; "
        "restoring an index file from a different point in time than the database is a common "
        "cause of LUM-1002 errors.",
    ),
    ("pb", None),
    ("h", "11. Recommended defaults"),
    (
        "p",
        "The following defaults apply to every new RAG feature. They are the starting point "
        "for experiments run by the retrieval optimisation process; changes are accepted only "
        "with an evaluation report.",
    ),
    (
        "t",
        [
            ["Parameter", "Team default", "Notes"],
            ["Chunking strategy", "Recursive, structure-aware", "Split at headings first"],
            ["Chunk size", "800 characters", "About 150-200 tokens of English text"],
            ["Chunk overlap", "100 characters", "12.5% of chunk size"],
            ["Tables", "Never split mid-row", "Repeat header row in every chunk"],
            ["Embedding dimension", "768", "Full dimension allowed in experiments"],
            ["Similarity", "Cosine", "Vectors L2-normalised, inner-product index"],
            ["Vector index", "Flat (exact)", "HNSW above 1 million chunks"],
            ["HNSW parameters", "M = 32, efSearch = 128", "Only when HNSW is used"],
            ["Hybrid search", "Enabled", "Dense + BM25, 50 candidates each"],
            ["Fusion", "RRF, k = 60", "Rank-based, no score normalisation"],
            ["Reranker", "Cross-encoder on top 50", "Keeps top 5"],
            ["Final top_k", "5", "Passages sent to the generator"],
            ["Abstention threshold", "Calibrated per model", "Never copied between models"],
            ["Max context", "6,000 tokens", "Retrieved passages only"],
            ["Generator temperature", "0.2", "As in the AI Engineering Handbook"],
        ],
    ),
    ("pb", None),
    ("h", "12. Anti-patterns and review checklist"),
    (
        "l",
        [
            "Tuning prompts before measuring retrieval. Check Recall@5 first.",
            "Copying a similarity threshold from another embedding model.",
            "Mixing vectors from two embedding models in one index.",
            "Treating the vector index as the source of truth and losing chunk offsets.",
            "Splitting tables so that rows lose their headers.",
            "Evaluating with chunk ids as relevance labels; labels must be evidence spans so "
            "that metrics stay comparable when chunking changes.",
            "Letting instructions found inside retrieved documents change the system's behaviour.",
            "Shipping without abstention: a RAG system that always answers will hallucinate on "
            "out-of-scope questions.",
        ],
    ),
    (
        "p",
        "Before a RAG feature is released, the Retrieval squad tech lead reviews the "
        "configuration against this guide, and the evaluation report must pass every blocking "
        "release gate in the LLM Evaluation Playbook.",
    ),
]

# --------------------------------------------------------------------------------------------
# llm_evaluation_playbook.pdf
# --------------------------------------------------------------------------------------------

EVAL_PLAYBOOK = [
    (
        "p",
        "LLM Evaluation Playbook - version 1.4 - effective 2026-07-01. Owner: Tomasz "
        "Wierzbicki, Evaluation Lead.",
    ),
    ("h", "1. Why we evaluate"),
    (
        "p",
        "LLM features fail quietly: a prompt tweak, a new chunk size or a model upgrade can "
        "lower answer quality without raising a single error. Every change to a production "
        "LLM feature is therefore measured on a fixed golden set before release, and the "
        "results are compared with the current production version.",
    ),
    ("h", "2. Golden sets"),
    (
        "l",
        [
            "Each production assistant has a golden set of at least 150 questions.",
            "At least 15% of the questions must be unanswerable from the corpus, to measure "
            "abstention.",
            "Relevance is labelled by evidence span: the document and the character offsets of "
            "the text that supports the answer. Chunk ids are never used as labels, because "
            "they change whenever chunking changes.",
            "Each question also has a reference answer written by a domain expert.",
            "Golden sets are refreshed every quarter and after every AI-SEV1 or AI-SEV2 incident.",
        ],
    ),
    ("h", "3. Retrieval metrics"),
    (
        "p",
        "A retrieved chunk counts as relevant when its character span overlaps an evidence "
        "span of the question. Retrieval metrics are computed on answerable questions only, at "
        "a cut-off K equal to the number of passages sent to the generator.",
    ),
    (
        "l",
        [
            "Precision@K: the fraction of the top K retrieved chunks that are relevant.",
            "Recall@K: the fraction of the question's evidence spans covered by at least one of "
            "the top K chunks.",
            "Hit rate@K: the fraction of questions with at least one relevant chunk in the top K.",
            "MRR (mean reciprocal rank): the average over questions of 1 / rank of the first "
            "relevant chunk, counting 0 when no relevant chunk is retrieved. MRR rewards putting "
            "the right passage first.",
            "nDCG@K is reported for information but is not a release gate.",
        ],
    ),
    ("h", "4. LLM-as-judge"),
    (
        "p",
        "Answer quality is scored by an LLM judge, a Tier-1 model run at temperature 0 with a "
        "fixed rubric. The judge scores three dimensions, each normalised to the range 0 to 1:",
    ),
    (
        "l",
        [
            "Faithfulness: every claim in the answer is supported by the retrieved passages.",
            "Answer relevance: the answer addresses the question that was asked.",
            "Correctness: the answer agrees with the reference answer.",
        ],
    ),
    (
        "p",
        "LLM judges have known biases: position bias (preferring the first of two options), "
        "verbosity bias (preferring longer answers) and self-preference (rating output from "
        "their own model family higher). The team mitigates them by scoring one answer at a "
        "time against a rubric instead of comparing pairs, by asking for a short justification "
        "before the score, and by calibrating the judge against human labels. A judge prompt "
        "may be used only if it agrees with human labels on at least 85% of a 50-question "
        "calibration sample.",
    ),
    ("h", "5. Abstention accuracy"),
    (
        "p",
        "Abstention accuracy is the fraction of questions where the system correctly abstained "
        "on an unanswerable question or correctly answered an answerable one. It is reported "
        "together with the false-abstention rate (answerable questions the system refused) "
        "because a system can trivially avoid hallucinations by refusing everything.",
    ),
    ("h", "6. Release gates"),
    (
        "p",
        "A release candidate must meet every blocking threshold. In addition, no blocking "
        "metric may drop by more than 2 points (0.02) compared with the current production "
        "version, even if it is still above its threshold.",
    ),
    (
        "t",
        [
            ["Metric", "Threshold", "Blocking"],
            ["Recall@5", ">= 0.85", "Yes"],
            ["MRR", ">= 0.70", "Yes"],
            ["Faithfulness", ">= 0.90", "Yes"],
            ["Answer correctness", ">= 0.80", "Yes"],
            ["Abstention accuracy", ">= 0.90", "Yes"],
            ["Answer relevance", ">= 0.85", "No (warning)"],
            ["P95 latency, fast mode", "<= 4 seconds", "No (warning)"],
            ["P95 latency, agentic mode", "<= 12 seconds", "No (warning)"],
            ["Cost per 1K requests", "Within Model Selection Policy cap", "Yes"],
        ],
    ),
    (
        "p",
        "The CI pipeline runs the offline evaluation on every pull request that changes "
        "prompts or retrieval configuration and fails the build when a blocking gate fails.",
    ),
    ("h", "7. Online monitoring"),
    (
        "p",
        "In production, 2% of requests are sampled and scored by the judge for faithfulness. "
        "The hallucination rate is one minus the faithfulness pass rate on this sample. A "
        "hallucination rate above 8% over a rolling 24-hour window raises an AI-SEV3 incident, "
        "and above 15% an AI-SEV2 incident, as described in the LLMOps Incident Runbook.",
    ),
    ("h", "8. Comparing runs"),
    (
        "p",
        "Every evaluation run stores the full configuration it used (chunking, retrieval, "
        "model tiers, prompt versions) together with its metrics, so that any two runs can be "
        "compared and any result can be reproduced. A comparison report lists each metric for "
        "both runs, the difference, and the questions whose outcome changed.",
    ),
]


# --------------------------------------------------------------------------------------------
# Writers
# --------------------------------------------------------------------------------------------


def write_docx(path: Path, title: str, blocks: list) -> None:
    doc = Document()
    doc.add_heading(title, level=0)
    for kind, content in blocks:
        if kind == "h":
            doc.add_heading(content, level=1)
        elif kind == "p":
            doc.add_paragraph(content)
        elif kind == "l":
            for item in content:
                doc.add_paragraph(item, style="List Bullet")
        elif kind == "pb":
            doc.add_page_break()
        else:
            table = doc.add_table(rows=len(content), cols=len(content[0]))
            table.style = "Table Grid"
            for r, row in enumerate(content):
                for c, cell in enumerate(row):
                    table.cell(r, c).text = cell
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


def write_pdf(path: Path, title: str, blocks: list) -> None:
    styles = getSampleStyleSheet()
    cell = styles["BodyText"]
    story = [Paragraph(escape(title), styles["Title"])]
    for kind, content in blocks:
        if kind == "h":
            story.append(Paragraph(escape(content), styles["Heading2"]))
        elif kind == "p":
            story.append(Paragraph(escape(content), styles["BodyText"]))
        elif kind == "l":
            story += [Paragraph("- " + escape(item), styles["BodyText"]) for item in content]
        elif kind == "pb":
            story.append(PageBreak())
            continue
        else:
            rows = [[Paragraph(escape(c), cell) for c in row] for row in content]
            table = Table(rows, repeatRows=1)
            table.setStyle(
                TableStyle(
                    [
                        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                        ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ]
                )
            )
            story.append(table)
        story.append(Spacer(1, 8))
    path.parent.mkdir(parents=True, exist_ok=True)
    SimpleDocTemplate(str(path), pagesize=A4, title=title).build(story)


def main() -> None:
    outputs = [
        (
            write_docx,
            CORPUS / "ai_engineering_handbook.docx",
            "Lumora Labs AI Engineering Handbook",
            HANDBOOK,
        ),
        (
            write_docx,
            CORPUS / "model_selection_policy_v1.docx",
            "Lumora Labs Model Selection Policy",
            MODEL_POLICY_V1,
        ),
        (write_pdf, CORPUS / "rag_design_guide.pdf", "Lumora Labs RAG Design Guide", RAG_GUIDE),
        (
            write_pdf,
            CORPUS / "llm_evaluation_playbook.pdf",
            "Lumora Labs LLM Evaluation Playbook",
            EVAL_PLAYBOOK,
        ),
        (
            write_docx,
            UPDATES / "model_selection_policy_v1.docx",
            "Lumora Labs Model Selection Policy",
            MODEL_POLICY_V2,
        ),
    ]
    for writer, path, title, blocks in outputs:
        writer(path, title, blocks)
        print(f"wrote {path.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
